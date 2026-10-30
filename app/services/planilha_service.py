"""Modelo padrão da planilha Excel e leitura/validação do arquivo enviado pelo usuário.

Este módulo não acessa o banco: recebe bytes e devolve dados validados (ou uma lista de erros).

Regra de impostos (card B03): vendas e recebíveis são informados em valor BRUTO.
O sistema estima o tributo uma única vez, nos cálculos, para nunca descontar imposto duas vezes.
"""

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.erros import PlanilhaInvalidaError

ABA_INSTRUCOES = "Instruções"
ABA_SALDO = "Saldo inicial"
ABA_VENDAS = "Vendas"
ABA_RECEBIVEIS = "Recebíveis"
ABA_DESPESAS = "Despesas"

TIPOS_VENDA = ("Pix", "Cartão de crédito", "Cartão de débito", "Boleto", "Dinheiro", "Outro")

# (cabeçalho, obrigatório). A ordem é a do modelo, mas na leitura a ordem das colunas não importa.
COLUNAS: dict[str, list[tuple[str, bool]]] = {
    ABA_SALDO: [("Data de referência", True), ("Saldo disponível (R$)", True)],
    ABA_VENDAS: [("Data", True), ("Descrição", True), ("Valor bruto (R$)", True), ("Tipo de venda", True)],
    ABA_RECEBIVEIS: [("Data prevista", True), ("Valor esperado (R$)", True), ("Descrição", False)],
    ABA_DESPESAS: [("Descrição", True), ("Vencimento", True), ("Valor (R$)", True)],
}

MAX_ERROS = 200
MAX_TEXTO = 200
DATA_MINIMA = date(2000, 1, 1)
DATA_MAXIMA = date(2100, 12, 31)
LIMITE_VALOR = Decimal("1000000000000")  # cabe em NUMERIC(14, 2)
CENTAVOS = Decimal("0.01")


@dataclass(frozen=True)
class SaldoDados:
    data_referencia: date
    saldo: Decimal


@dataclass(frozen=True)
class VendaDados:
    data: date
    descricao: str
    valor_bruto: Decimal
    tipo: str


@dataclass(frozen=True)
class RecebivelDados:
    data_prevista: date
    valor_esperado: Decimal
    descricao: str | None


@dataclass(frozen=True)
class DespesaDados:
    descricao: str
    vencimento: date
    valor: Decimal


@dataclass
class DadosPlanilha:
    saldo: SaldoDados
    vendas: list[VendaDados] = field(default_factory=list)
    recebiveis: list[RecebivelDados] = field(default_factory=list)
    despesas: list[DespesaDados] = field(default_factory=list)


# ---------------------------------------------------------------- modelo (download)

def gerar_modelo(max_linhas: int = 5000, max_mb: int = 5) -> bytes:
    """Monta a planilha modelo (.xlsx) com as abas, cabeçalhos e formatos definidos."""
    wb = Workbook()
    ws = wb.active
    ws.title = ABA_INSTRUCOES
    _preencher_instrucoes(ws, max_linhas, max_mb)

    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="14213D")
    formatos = {
        ABA_SALDO: {"A": "dd/mm/yyyy", "B": "#,##0.00"},
        ABA_VENDAS: {"A": "dd/mm/yyyy", "C": "#,##0.00"},
        ABA_RECEBIVEIS: {"A": "dd/mm/yyyy", "B": "#,##0.00"},
        ABA_DESPESAS: {"B": "dd/mm/yyyy", "C": "#,##0.00"},
    }
    larguras = {ABA_SALDO: (22, 24), ABA_VENDAS: (14, 40, 20, 20), ABA_RECEBIVEIS: (16, 22, 40), ABA_DESPESAS: (40, 16, 18)}

    for nome_aba, colunas in COLUNAS.items():
        aba = wb.create_sheet(nome_aba)
        for i, (nome, _) in enumerate(colunas, start=1):
            celula = aba.cell(row=1, column=i, value=nome)
            celula.font = cabecalho_fonte
            celula.fill = cabecalho_fundo
            celula.alignment = Alignment(horizontal="center", vertical="center")
            aba.column_dimensions[get_column_letter(i)].width = larguras[nome_aba][i - 1]
        aba.freeze_panes = "A2"
        ultima = 2 if nome_aba == ABA_SALDO else 1001
        for letra, formato in formatos[nome_aba].items():
            for linha in range(2, ultima + 1):
                aba[f"{letra}{linha}"].number_format = formato

    validacao = DataValidation(
        type="list", formula1='"' + ",".join(TIPOS_VENDA) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="Tipo de venda", error="Escolha um dos tipos da lista.",
    )
    wb[ABA_VENDAS].add_data_validation(validacao)
    validacao.add("D2:D1001")

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _preencher_instrucoes(ws, max_linhas: int, max_mb: int) -> None:
    linhas = [
        ("SplitAgent — Planilha modelo", "titulo"),
        ("", None),
        ("Como preencher", "secao"),
        ("1. Preencha as abas Saldo inicial, Vendas, Recebíveis e Despesas. Não altere os nomes das abas nem os cabeçalhos da linha 1.", None),
        ("2. Datas no formato dd/mm/aaaa. Valores em reais, com até 2 casas decimais (ex.: 1234,56), sem texto.", None),
        ("3. Linhas totalmente vazias são ignoradas. Toda linha preenchida precisa ter todos os campos obrigatórios.", None),
        (f"4. Envie um arquivo .xlsx de até {max_mb} MB, com no máximo {max_linhas} linhas preenchidas por aba.", None),
        ("", None),
        ("O que vai em cada aba", "secao"),
        ("Saldo inicial: UMA linha com a data de referência e o saldo disponível em caixa nessa data (pode ser negativo).", None),
        ("Vendas: uma linha por venda já realizada, com data, descrição, valor BRUTO e tipo de venda (Pix, Cartão de crédito, Cartão de débito, Boleto, Dinheiro ou Outro).", None),
        ("Recebíveis: valores que você ainda vai receber (por exemplo, parcelas no cartão), com a data prevista e o valor esperado. A descrição é opcional.", None),
        ("Despesas: contas a pagar, com descrição, data de vencimento e valor.", None),
        ("", None),
        ("Impostos: informe sempre valores BRUTOS", "secao"),
        ("Em Vendas e Recebíveis, informe o valor cheio, ANTES de descontar qualquer imposto (CBS/IBS). O SplitAgent estima o tributo uma única vez.", None),
        ("Se você já descontar o imposto na planilha, ele será descontado duas vezes e a previsão do caixa ficará errada.", None),
        ("", None),
        ("Exemplos de linha", "secao"),
        ("Vendas: 15/11/2026 | Venda de balcão | 250,00 | Pix", None),
        ("Recebíveis: 20/11/2026 | 1.200,00 | Parcela 2 de 3 — cartão", None),
        ("Despesas: Aluguel | 05/12/2026 | 1.800,00", None),
    ]
    ws.column_dimensions["A"].width = 120
    for i, (texto, estilo) in enumerate(linhas, start=1):
        celula = ws.cell(row=i, column=1, value=texto)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        if estilo == "titulo":
            celula.font = Font(bold=True, size=16)
        elif estilo == "secao":
            celula.font = Font(bold=True, size=12, color="14213D")


# ---------------------------------------------------------------- leitura e validação

def _norm(texto) -> str:
    """Compara nomes sem acento, sem maiúsculas e sem o trecho entre parênteses, ex.: (R$)."""
    s = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    s = re.sub(r"\(.*?\)", "", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def _vazio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


class _Erros:
    def __init__(self) -> None:
        self.itens: list[dict] = []
        self.truncado = False

    def __bool__(self) -> bool:
        return bool(self.itens)

    def add(self, aba: str, linha: int | None, coluna: str | None, mensagem: str) -> None:
        if len(self.itens) >= MAX_ERROS:
            self.truncado = True
            return
        self.itens.append({"aba": aba, "linha": linha, "coluna": coluna, "mensagem": mensagem})


def _ler_data(valor, erros: _Erros, aba: str, linha: int, coluna: str) -> date | None:
    if _vazio(valor):
        erros.add(aba, linha, coluna, "Campo obrigatório vazio.")
        return None
    resultado = None
    if isinstance(valor, datetime):
        resultado = valor.date()
    elif isinstance(valor, date):
        resultado = valor
    elif isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                resultado = datetime.strptime(valor.strip(), formato).date()
                break
            except ValueError:
                continue
    if resultado is None:
        erros.add(aba, linha, coluna, "Data inválida. Use o formato dd/mm/aaaa.")
        return None
    if not DATA_MINIMA <= resultado <= DATA_MAXIMA:
        erros.add(aba, linha, coluna, "Data fora do intervalo aceito (2000 a 2100). Confira o ano.")
        return None
    return resultado


def _ler_dinheiro(valor, erros: _Erros, aba: str, linha: int, coluna: str, permitir_negativo: bool = False) -> Decimal | None:
    if _vazio(valor):
        erros.add(aba, linha, coluna, "Campo obrigatório vazio.")
        return None
    mensagem_invalido = "Valor numérico inválido. Use apenas números, ex.: 1234,56."
    try:
        if isinstance(valor, bool):
            raise InvalidOperation
        if isinstance(valor, Decimal):
            numero = valor
        elif isinstance(valor, (int, float)):
            numero = Decimal(str(valor))
        elif isinstance(valor, str):
            texto = valor.strip().replace("R$", "").replace("\xa0", "").replace(" ", "")
            if "," in texto:
                texto = texto.replace(".", "").replace(",", ".")
            numero = Decimal(texto)
        else:
            raise InvalidOperation
        if not numero.is_finite():
            raise InvalidOperation
        numero = numero.quantize(CENTAVOS, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        erros.add(aba, linha, coluna, mensagem_invalido)
        return None
    if abs(numero) >= LIMITE_VALOR:
        erros.add(aba, linha, coluna, "Valor muito grande.")
        return None
    if not permitir_negativo and numero <= 0:
        erros.add(aba, linha, coluna, "O valor deve ser maior que zero.")
        return None
    return numero


def _ler_texto(valor, erros: _Erros, aba: str, linha: int, coluna: str, obrigatorio: bool = True) -> str | None:
    if _vazio(valor):
        if obrigatorio:
            erros.add(aba, linha, coluna, "Campo obrigatório vazio.")
        return None
    texto = str(valor).strip()
    if len(texto) > MAX_TEXTO:
        erros.add(aba, linha, coluna, f"Texto muito longo (máximo de {MAX_TEXTO} caracteres).")
        return None
    return texto


def _ler_tipo_venda(valor, erros: _Erros, aba: str, linha: int, coluna: str) -> str | None:
    texto = _ler_texto(valor, erros, aba, linha, coluna)
    if texto is None:
        return None
    for tipo in TIPOS_VENDA:
        if _norm(tipo) == _norm(texto):
            return tipo
    erros.add(aba, linha, coluna, f"Tipo de venda inválido. Use um destes: {', '.join(TIPOS_VENDA)}.")
    return None


def _linhas_da_aba(abas: dict, nome_aba: str, erros: _Erros, max_linhas: int) -> list[tuple[int, dict]] | None:
    """Devolve [(número da linha no Excel, {cabeçalho: valor})]. None se a estrutura da aba estiver errada."""
    ws = abas.get(_norm(nome_aba))
    if ws is None:
        erros.add(nome_aba, None, None, f'Aba "{nome_aba}" não encontrada. Use a planilha modelo da plataforma.')
        return None
    iterador = ws.iter_rows(values_only=True)
    cabecalho = next(iterador, None)
    if cabecalho is None or all(_vazio(c) for c in cabecalho):
        erros.add(nome_aba, 1, None, "Aba vazia: a linha 1 deve conter os cabeçalhos do modelo.")
        return None
    posicoes: dict[str, int] = {}
    for indice, nome in enumerate(cabecalho):
        if not _vazio(nome):
            posicoes.setdefault(_norm(nome), indice)
    faltando = [nome for nome, obrigatoria in COLUNAS[nome_aba] if obrigatoria and _norm(nome) not in posicoes]
    for nome in faltando:
        erros.add(nome_aba, 1, nome, f'Coluna "{nome}" não encontrada na linha 1.')
    if faltando:
        return None
    linhas: list[tuple[int, dict]] = []
    for numero, valores in enumerate(iterador, start=2):
        if all(_vazio(v) for v in valores):
            continue
        if len(linhas) >= max_linhas:
            erros.add(nome_aba, numero, None, f"A aba tem mais de {max_linhas} linhas preenchidas. Remova linhas antigas ou divida os dados.")
            break
        registro = {}
        for nome, _ in COLUNAS[nome_aba]:
            i = posicoes.get(_norm(nome))
            registro[nome] = valores[i] if i is not None and i < len(valores) else None
        linhas.append((numero, registro))
    return linhas


def ler_planilha(conteudo: bytes, max_linhas: int) -> DadosPlanilha:
    """Lê e valida a planilha. Em caso de qualquer problema, levanta PlanilhaInvalidaError com todos os erros."""
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:  # arquivo corrompido, não é xlsx, etc.
        raise PlanilhaInvalidaError(
            "Não foi possível ler o arquivo. Envie uma planilha Excel (.xlsx) válida, baseada no modelo da plataforma.",
            [{"aba": "Arquivo", "linha": None, "coluna": None, "mensagem": "Arquivo ilegível ou corrompido."}],
        ) from exc

    erros = _Erros()
    try:
        abas = {_norm(nome): wb[nome] for nome in wb.sheetnames}
        saldo = _ler_saldo(abas, erros, max_linhas)
        vendas = _ler_vendas(abas, erros, max_linhas)
        recebiveis = _ler_recebiveis(abas, erros, max_linhas)
        despesas = _ler_despesas(abas, erros, max_linhas)
    finally:
        wb.close()

    if not erros and not (vendas or recebiveis or despesas):
        erros.add("Planilha", None, None, "Preencha ao menos uma venda, um recebível ou uma despesa.")
    if erros:
        quantidade = len(erros.itens)
        raise PlanilhaInvalidaError(
            f"Encontramos {quantidade} problema{'s' if quantidade != 1 else ''} na planilha. Corrija e envie novamente.",
            erros.itens,
            erros.truncado,
        )
    return DadosPlanilha(saldo=saldo, vendas=vendas, recebiveis=recebiveis, despesas=despesas)


def _ler_saldo(abas, erros: _Erros, max_linhas: int) -> SaldoDados | None:
    linhas = _linhas_da_aba(abas, ABA_SALDO, erros, max_linhas)
    if linhas is None:
        return None
    if not linhas:
        erros.add(ABA_SALDO, 2, None, "Informe o saldo inicial na linha 2.")
        return None
    if len(linhas) > 1:
        erros.add(ABA_SALDO, linhas[1][0], None, "Informe apenas uma linha de saldo inicial.")
    numero, registro = linhas[0]
    data = _ler_data(registro["Data de referência"], erros, ABA_SALDO, numero, "Data de referência")
    saldo = _ler_dinheiro(registro["Saldo disponível (R$)"], erros, ABA_SALDO, numero, "Saldo disponível (R$)", permitir_negativo=True)
    if data is None or saldo is None:
        return None
    return SaldoDados(data, saldo)


def _ler_vendas(abas, erros: _Erros, max_linhas: int) -> list[VendaDados]:
    resultado = []
    for numero, r in _linhas_da_aba(abas, ABA_VENDAS, erros, max_linhas) or []:
        data = _ler_data(r["Data"], erros, ABA_VENDAS, numero, "Data")
        descricao = _ler_texto(r["Descrição"], erros, ABA_VENDAS, numero, "Descrição")
        valor = _ler_dinheiro(r["Valor bruto (R$)"], erros, ABA_VENDAS, numero, "Valor bruto (R$)")
        tipo = _ler_tipo_venda(r["Tipo de venda"], erros, ABA_VENDAS, numero, "Tipo de venda")
        if None not in (data, descricao, valor, tipo):
            resultado.append(VendaDados(data, descricao, valor, tipo))
    return resultado


def _ler_recebiveis(abas, erros: _Erros, max_linhas: int) -> list[RecebivelDados]:
    resultado = []
    for numero, r in _linhas_da_aba(abas, ABA_RECEBIVEIS, erros, max_linhas) or []:
        data = _ler_data(r["Data prevista"], erros, ABA_RECEBIVEIS, numero, "Data prevista")
        valor = _ler_dinheiro(r["Valor esperado (R$)"], erros, ABA_RECEBIVEIS, numero, "Valor esperado (R$)")
        descricao = _ler_texto(r["Descrição"], erros, ABA_RECEBIVEIS, numero, "Descrição", obrigatorio=False)
        if None not in (data, valor):
            resultado.append(RecebivelDados(data, valor, descricao))
    return resultado


def _ler_despesas(abas, erros: _Erros, max_linhas: int) -> list[DespesaDados]:
    resultado = []
    for numero, r in _linhas_da_aba(abas, ABA_DESPESAS, erros, max_linhas) or []:
        descricao = _ler_texto(r["Descrição"], erros, ABA_DESPESAS, numero, "Descrição")
        vencimento = _ler_data(r["Vencimento"], erros, ABA_DESPESAS, numero, "Vencimento")
        valor = _ler_dinheiro(r["Valor (R$)"], erros, ABA_DESPESAS, numero, "Valor (R$)")
        if None not in (descricao, vencimento, valor):
            resultado.append(DespesaDados(descricao, vencimento, valor))
    return resultado
