from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.core import config, database
from app.services import importacao_service, movimentacao_service
from tests.planilhas import XLSX, montar_planilha

VENDAS = [
    (date(2026, 10, 1), "Venda balcão", 250, "Pix"),
    (date(2026, 10, 2), "Venda online", 1000.5, "Cartão de crédito"),
]
RECEBIVEIS = [(date(2026, 11, 10), 1200, "Parcela 2 de 3"), (date(2026, 11, 20), 800, None)]
DESPESAS = [("Aluguel", date(2026, 11, 5), 1800)]


def planilha_valida(**extra) -> bytes:
    return montar_planilha(vendas=VENDAS, recebiveis=RECEBIVEIS, despesas=DESPESAS, **extra)


def enviar(client, conteudo, nome="planilha.xlsx", **params):
    return client.post("/api/v1/importacoes", files={"arquivo": (nome, conteudo, XLSX)}, params=params)


def local(erro):
    return (erro["aba"], erro["linha"], erro["coluna"])


# ---------------------------------------------------------------- modelo

def test_modelo_tem_abas_e_cabecalhos(client_db):
    resposta = client_db.get("/api/v1/importacoes/modelo")
    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith(XLSX)
    assert "modelo_splitagent.xlsx" in resposta.headers["content-disposition"]
    wb = load_workbook(BytesIO(resposta.content))
    assert wb.sheetnames == ["Instruções", "Saldo inicial", "Vendas", "Recebíveis", "Despesas"]
    assert [c.value for c in wb["Vendas"][1]] == ["Data", "Descrição", "Valor bruto (R$)", "Tipo de venda"]
    assert [c.value for c in wb["Despesas"][1]] == ["Descrição", "Vencimento", "Valor (R$)"]


def test_modelo_nao_exige_login(client):
    assert client.get("/api/v1/importacoes/modelo").status_code == 200


def test_modelo_em_branco_nao_importa(client_db):
    modelo = client_db.get("/api/v1/importacoes/modelo").content
    resposta = enviar(client_db, modelo)
    assert resposta.status_code == 422
    assert any(e["aba"] == "Saldo inicial" for e in resposta.json()["erros"])


# ---------------------------------------------------------------- importação válida

def test_importa_planilha_valida(client_db):
    resposta = enviar(client_db, planilha_valida())
    assert resposta.status_code == 201
    corpo = resposta.json()
    assert corpo["nome_arquivo"] == "planilha.xlsx"
    assert corpo["total_vendas"] == 2
    assert corpo["total_recebiveis"] == 2
    assert corpo["total_despesas"] == 1
    assert corpo["substituiu_dados_anteriores"] is False
    assert corpo["saldo_inicial"] == {"data_referencia": "2026-11-01", "saldo": "5000.00"}

    atual = client_db.get("/api/v1/importacoes/atual")
    assert atual.status_code == 200
    assert atual.json()["importacao_id"] == corpo["importacao_id"]


def test_atual_sem_importacao_retorna_404(client_db):
    resposta = client_db.get("/api/v1/importacoes/atual")
    assert resposta.status_code == 404
    assert "Nenhuma planilha" in resposta.json()["mensagem"]


def test_valores_em_texto_ptbr_e_saldo_negativo(client_db):
    conteudo = montar_planilha(saldo=("01/11/2026", "-R$ 250,10"), despesas=[("Luz", "05/12/2026", "1.234,56")])
    resposta = enviar(client_db, conteudo)
    assert resposta.status_code == 201
    assert resposta.json()["saldo_inicial"] == {"data_referencia": "2026-11-01", "saldo": "-250.10"}


def test_cabecalhos_sem_acento_e_em_minusculas_sao_aceitos(client_db):
    conteudo = montar_planilha(
        vendas=VENDAS,
        cabecalhos={"Vendas": ["data", "descricao", "VALOR BRUTO", "tipo de venda"]},
    )
    assert enviar(client_db, conteudo).status_code == 201


# ---------------------------------------------------------------- arquivo inválido

def test_extensao_diferente_de_xlsx_retorna_415(client_db):
    resposta = enviar(client_db, b"a,b,c", nome="dados.csv")
    assert resposta.status_code == 415
    assert ".xlsx" in resposta.json()["mensagem"]


def test_arquivo_que_nao_e_zip_retorna_415(client_db):
    assert enviar(client_db, b"isto nao e excel").status_code == 415


def test_xlsx_corrompido_retorna_422_com_mensagem_clara(client_db):
    resposta = enviar(client_db, b"PK\x03\x04lixo")
    assert resposta.status_code == 422
    assert "Não foi possível ler" in resposta.json()["mensagem"]


def test_arquivo_acima_do_limite_retorna_413(client_db, monkeypatch):
    monkeypatch.setenv("IMPORTACAO_MAX_MB", "1")
    config.get_settings.cache_clear()
    resposta = enviar(client_db, b"PK" + b"0" * (1024 * 1024))
    assert resposta.status_code == 413


def test_sem_usuario_habilitado_retorna_501(client_db, monkeypatch):
    monkeypatch.setenv("USUARIO_DEV_ID", "")
    config.get_settings.cache_clear()
    assert enviar(client_db, planilha_valida()).status_code == 501
    assert client_db.get("/api/v1/importacoes/atual").status_code == 501
    assert client_db.get("/api/v1/movimentacoes").status_code == 501


# ---------------------------------------------------------------- erros de conteúdo

def test_reporta_todos_os_erros_e_nao_salva_nada(client_db):
    conteudo = montar_planilha(
        saldo=("31/02/2026", 100),
        vendas=[
            (date(2026, 10, 1), "", 100, "Pix"),
            (date(2026, 10, 2), "Ok", -5, "Pix"),
            (date(2026, 10, 3), "Ok", 10, "Cheque"),
            (date(2026, 10, 4), "Venda válida", 10, "Pix"),
        ],
        recebiveis=[(date(2026, 11, 1), "abc", None)],
        despesas=[("Luz", None, 50)],
    )
    resposta = enviar(client_db, conteudo)
    assert resposta.status_code == 422
    corpo = resposta.json()
    locais = {local(e) for e in corpo["erros"]}
    assert ("Saldo inicial", 2, "Data de referência") in locais
    assert ("Vendas", 2, "Descrição") in locais
    assert ("Vendas", 3, "Valor bruto (R$)") in locais
    assert ("Vendas", 4, "Tipo de venda") in locais
    assert ("Recebíveis", 2, "Valor esperado (R$)") in locais
    assert ("Despesas", 2, "Vencimento") in locais
    assert not any(e["linha"] == 5 for e in corpo["erros"])
    assert corpo["total_erros"] == len(corpo["erros"]) == 6
    assert corpo["erros_truncados"] is False
    assert client_db.get("/api/v1/importacoes/atual").status_code == 404


def test_coluna_e_aba_ausentes(client_db):
    conteudo = montar_planilha(
        vendas=VENDAS,
        omitir_abas=("Despesas",),
        cabecalhos={"Vendas": ["Data", "Descrição", "Valor bruto (R$)"]},
    )
    resposta = enviar(client_db, conteudo)
    assert resposta.status_code == 422
    locais = {local(e) for e in resposta.json()["erros"]}
    assert ("Vendas", 1, "Tipo de venda") in locais
    assert ("Despesas", None, None) in locais


def test_saldo_com_mais_de_uma_linha_e_rejeitado(client_db):
    conteudo = montar_planilha(vendas=VENDAS)
    wb = load_workbook(BytesIO(conteudo))
    wb["Saldo inicial"].append([date(2026, 11, 2), 10])
    saida = BytesIO()
    wb.save(saida)
    resposta = enviar(client_db, saida.getvalue())
    assert resposta.status_code == 422
    assert ("Saldo inicial", 3, None) in {local(e) for e in resposta.json()["erros"]}


def test_limite_de_linhas_por_aba(client_db, monkeypatch):
    monkeypatch.setenv("IMPORTACAO_MAX_LINHAS", "2")
    config.get_settings.cache_clear()
    conteudo = montar_planilha(vendas=[(date(2026, 10, i), f"V{i}", 10, "Pix") for i in range(1, 6)])
    resposta = enviar(client_db, conteudo)
    assert resposta.status_code == 422
    assert "mais de 2 linhas" in resposta.json()["erros"][0]["mensagem"]


# ---------------------------------------------------------------- atualização da planilha

def test_segunda_importacao_exige_confirmacao_para_substituir(client_db):
    primeira = enviar(client_db, planilha_valida()).json()
    bloqueada = enviar(client_db, montar_planilha(despesas=[("Nova", date(2026, 12, 1), 10)]))
    assert bloqueada.status_code == 409
    assert client_db.get("/api/v1/importacoes/atual").json()["importacao_id"] == primeira["importacao_id"]

    nova = enviar(client_db, montar_planilha(despesas=[("Nova", date(2026, 12, 1), 10)]), substituir="true")
    assert nova.status_code == 201
    corpo = nova.json()
    assert corpo["substituiu_dados_anteriores"] is True
    assert (corpo["total_vendas"], corpo["total_recebiveis"], corpo["total_despesas"]) == (0, 0, 1)
    atual = client_db.get("/api/v1/importacoes/atual").json()
    assert atual["importacao_id"] == corpo["importacao_id"]
    movimentacoes = client_db.get("/api/v1/movimentacoes", params={"data_inicio": "2026-01-01"}).json()
    assert movimentacoes["total_itens"] == 1  # só a despesa nova: os dados antigos foram trocados


def test_planilha_invalida_nao_apaga_dados_existentes(client_db):
    primeira = enviar(client_db, planilha_valida()).json()
    resposta = enviar(client_db, montar_planilha(saldo=None, vendas=VENDAS), substituir="true")
    assert resposta.status_code == 422
    atual = client_db.get("/api/v1/importacoes/atual").json()
    assert atual["importacao_id"] == primeira["importacao_id"]
    assert atual["total_vendas"] == 2


# ---------------------------------------------------------------- isolamento entre usuários

def test_dados_de_um_usuario_nao_aparecem_para_outro(client_db):
    sessao = database.get_session_factory()()
    try:
        importacao_service.importar(sessao, 1, "a.xlsx", montar_planilha(despesas=[("Do usuário 1", date(2030, 1, 1), 10)]), False)
        importacao_service.importar(sessao, 2, "b.xlsx", montar_planilha(despesas=[("Do usuário 2", date(2030, 1, 1), 20)]), False)
        do_1 = movimentacao_service.listar(sessao, 1, date(2030, 1, 1), None, None, 100, 0)
        do_2 = movimentacao_service.listar(sessao, 2, date(2030, 1, 1), None, None, 100, 0)
        assert [m["descricao"] for m in do_1["movimentacoes"]] == ["Do usuário 1"]
        assert [m["descricao"] for m in do_2["movimentacoes"]] == ["Do usuário 2"]
        assert importacao_service.obter_resumo_atual(sessao, 1)["nome_arquivo"] == "a.xlsx"
        assert importacao_service.obter_resumo_atual(sessao, 2)["nome_arquivo"] == "b.xlsx"
    finally:
        sessao.close()
